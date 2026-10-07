"""Build the hand-authored Word conversion sample and prepare Core payloads."""
import argparse
import hashlib
import json
import subprocess
import time
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

parser = argparse.ArgumentParser()
parser.add_argument("--core-cli", required=True)
parser.add_argument("--output-dir", required=True)
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
fixture = root / "apps/native-office/LaTeXSnipper.Word.HostTests/fixtures/word-standard-mixed-v1.json"
contract = json.loads(fixture.read_text(encoding="utf-8"))
output = Path(args.output_dir).resolve()
output.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.color.rgb = RGBColor(0, 0, 0)
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(9)
doc.styles["Title"].font.size = Pt(22)
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 2"].font.size = Pt(12)
doc.core_properties.author = "LaTeXSnipper"
doc.core_properties.title = "Word 公式转换标准样本文档"
doc.core_properties.subject = "固定样例 v1"
sections = ["基本公式与定界符", "矩阵分段与嵌套", "积分统计与重复公式", "容错与支持边界"]
descriptions = [
    "检查四种数学定界符，以及分数、根式和上下标。转换不得删除公式两侧的编号和结束标记。",
    "检查结构和内容是否完整。多行结构不能只留下源码、空公式或丢失矩阵元素。",
    "检查积分上下限、化学元素正体和重复公式。相同源码仍须具有独立对象标识。",
    "本页应保留原文。TikZ 应走绘图入口；错误输入不应覆盖正文；未闭合和裸源码不纳入普通批量扫描。",
]
timings = []
for page in range(4):
    if page:
        doc.add_page_break()
    else:
        doc.add_paragraph("Word 公式转换标准样本文档", "Title")
        doc.add_paragraph("固定样例 v1  共 16 项  2026 年 10 月")
        doc.add_paragraph("用于核对批量转换前后的结构、正文保留和可编辑性。公式内容逐条编写；本输入文档保留 LaTeX 源码，供转换功能处理。")
    doc.add_paragraph(sections[page], "Heading 1")
    doc.add_paragraph(descriptions[page])
    for case in contract["cases"][page * 4:page * 4 + 4]:
        doc.add_paragraph(case["name"] + " " + case["label"], "Heading 2")
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.add_run(case["name"] + " 正文前 ")
        run = paragraph.add_run(case["source"])
        run.font.name = "Consolas"
        run.font.size = Pt(11)
        paragraph.add_run(" 正文后 " + case["name"])
        if not case.get("preserve"):
            start = time.perf_counter()
            result = subprocess.run([args.core_cli, "render", "--latex", case["latex"], "--to", "omml"],
                                    capture_output=True, text=True, encoding="utf-8", timeout=30)
            elapsed = (time.perf_counter() - start) * 1000
            if result.returncode or "<m:oMath" not in result.stdout:
                raise RuntimeError(case["name"] + ": " + result.stderr)
            case["omml"] = result.stdout.strip()
            timings.append({"name": case["name"], "elapsedMs": elapsed})
        else:
            case["omml"] = ""
doc.save(output / "word-standard-mixed-input.docx")
contract["preparation"] = {
    "coreExecutableSha256": hashlib.sha256(Path(args.core_cli).read_bytes()).hexdigest(),
    "scope": "One CLI process per valid formula; process startup included, not the persistent production worker.",
    "timings": timings,
}
(output / "standard.generated.json").write_text(json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"cases": 16, "coreConverted": len(timings), "coreCliTotalMs": sum(x["elapsedMs"] for x in timings)}))
